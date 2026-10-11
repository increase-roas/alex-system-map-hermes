import json,hashlib,zipfile,sys,re
from pathlib import Path
from collections import Counter
import openpyxl
ROOT=Path(__file__).resolve().parents[2]
def load(n):return json.loads((ROOT/'data'/f'{n}.json').read_text())['records']
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
results=[]
def check(name,okay,detail):results.append(dict(item=name,status='VERIFIED' if okay else 'FAILED',evidence=detail))
cc=load('components');rr=load('relationships');ss=load('sources');ee=load('evidence');records=load('source-records');coverage=load('coverage');tables=load('table-index');layers=load('layers');issues=load('issues')
ci={x['id']:x for x in cc};ei={x['id']:x for x in ee};ri={x['id']:x for x in records};si={x['id']:x for x in ss}
w=openpyxl.load_workbook(ROOT/'inputs/Alex_Master_Build_and_Connectors.xlsx')
for name,rows in [('components',cc),('relationships',rr),('sources',ss),('evidence',ee),('records',records),('tables',tables),('layers',layers),('issues',issues)]:check(name+'_unique_ids',len({x['id'] for x in rows})==len(rows),str(len(rows))+' records inspected')
for file in ['Alex_Master_Build_and_Connectors.xlsx','Alex_Master_Source_Files.zip']:
 parent=ROOT.parent/file
 check('original_'+file+'_preserved',parent.exists() and sha(parent)==sha(ROOT/'inputs'/file),dict(source_path=str(parent),copy_path=str(ROOT/'inputs'/file),sha256=sha(ROOT/'inputs'/file)))
# Original workbook is the yardstick; do not accept a count copied from a builder report.
orig=[]
for row in range(5,w['MAP COVERAGE'].max_row+1):
 s=w['MAP COVERAGE'];a=s.cell(row,5).hyperlink;ptr=a.location or a.target.lstrip('#');sheet,pos=ptr.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'");n=int(re.search(r'\d+',pos).group())
 orig.append((row,s.cell(row,1).value,s.cell(row,2).value,sheet,n))
actual=[(x['coverage_row'],x['record_type'],ri[x['record_id']]['label'],x['source']['sheet'],x['source']['row']) for x in coverage]
check('original_2377_record_coverage_exact',len(orig)==len(records)==len(coverage)==2377 and sorted(orig)==sorted(actual) and Counter(x['record_id'] for x in coverage)==Counter(x['id'] for x in records),dict(original=len(orig),canonical=len(records),coverage=len(coverage),per_kind=dict(Counter(x[1] for x in orig))))
manifest=json.loads((ROOT/'inputs/coverage-audit-manifest.json').read_text())
maprows=manifest['components'];okay=len(maprows)==198
for x in maprows:
 match=[c for c in cc if c['name']==x['name'] and c['layer_id']==f"layer_{x['layer']:02}" and c['purpose']==x['purpose'] and c['boundary']==x['boundary']]
 okay &= bool(match) and w['MY SYSTEM MAP'].cell(x['map_row'],1).value==x['name']
check('198_map_components_preserved',okay,'Compared each canonical map identity/purpose/boundary/layer and original MY SYSTEM MAP row label')
check('eight_layers_preserved',len(layers)==8 and {x['id'] for x in layers}=={f'layer_{i:02}' for i in range(1,9)} and all(c['layer_id'] in {x['id'] for x in layers} for c in cc),[x['name'] for x in layers])
check('source_pointers_resolve',all(r['component_id'] in ci and r['placement_component_id'] in ci and all(x in ei for x in r['evidence_ids']) for r in records) and all(e['source_id'] in si and (not e.get('sheet') or e['sheet'] in w.sheetnames and 1<=e['row']<=w[e['sheet']].max_row) for e in ee) and all(all(x in ei for x in c['evidence_ids']) and all(x in ri for x in c['source_record_ids']) for c in cc),'All records/components/evidence source and row pointers checked')
archived=[s for s in ss if s['kind']=='archived_file'];bad=[]
with zipfile.ZipFile(ROOT/'inputs/Alex_Master_Source_Files.zip') as z:
 for s in archived:
  try:b=z.read(s['archive_member']);ok=hashlib.sha256(b).hexdigest()==s['sha256'] and len(b)==s['bytes']
  except Exception:ok=False
  if not ok:bad.append(s['id'])
check('1009_archived_file_bytes',len(archived)==1009 and not bad,dict(count=len(archived),bad_source_ids=bad))
# Exported originals may have recognized credentials redacted. Compare all other values exactly.
redaction_locations={x['location'] for x in json.loads((ROOT/'reports/redactions.json').read_text())['redactions']}
def expected_cell(c):
 d={'column':c.column_letter,'value':c.value}
 if c.comment:d['comment']=c.comment.text
 if c.hyperlink:d['link']=c.hyperlink.location or c.hyperlink.target
 return d
def matches(original,actual,loc):
 if original==actual:return True
 return isinstance(original,str) and isinstance(actual,str) and '[REDACTED credential-shaped value]' in actual and loc in redaction_locations
bad_tables=[];exported_cells=0
for t in tables:
 d=json.loads((ROOT/t['path']).read_text());s=w[t['sheet']];expectrows=[r for r in s if any(c.value is not None or c.comment for c in r)]
 if len(expectrows)!=len(d['rows']) or d['sheet']!=t['sheet']:bad_tables.append(t['id']);continue
 for er,ar in zip(expectrows,d['rows']):
  expect=[c for c in er if c.value is not None or c.comment]
  if ar['row']!=er[0].row or len(expect)!=len(ar['cells']):bad_tables.append(t['id']);continue
  for c,ac in zip(expect,ar['cells']):
   exported_cells+=1;ec=expected_cell(c);base=f'{s.title}!{c.coordinate}'
   if set(ec)!=set(ac) or not all(matches(ec[k],ac.get(k),base+(':'+k if k in ('comment','link') else '')) for k in ec):bad_tables.append(t['id'])
check('192_source_tables_exact_except_recorded_redactions',len(tables)==len(w.sheetnames)==192 and {x['sheet'] for x in tables}==set(w.sheetnames) and not bad_tables,dict(tables=len(tables),cells_compared=exported_cells,bad_tables=sorted(set(bad_tables))))
# Derive endpoint map using source rows, independently of edge count.
keymap={}
for r in records:
 if r.get('full_source'):
  p=r['full_source'];key=w[p['sheet']].cell(p['row'],15).value
  if key:keymap[key]=r['component_id']
relrows=w['S04-20 Relationships'];matchedrows=[];badrels=[]
for r in rr:
 evs=[ei.get(x) for x in r['evidence_ids']];evs=[x for x in evs if x]
 if len(evs)!=1 or evs[0].get('sheet')!='S04-20 Relationships':badrels.append(r['id']);continue
 n=evs[0]['row'];matchedrows.append(n);a,typ,b,label,ep,res=[relrows.cell(n,j).value for j in range(1,7)]
 okay=r['type']==str(typ) and r['label']==(label or str(typ)) and r['source_pointer']==ep and r['resolution']==res
 for key,endpoint in [(a,r['source_id']),(b,r['target_id'])]:
  okay &= endpoint in ci and (endpoint==keymap[key] if key in keymap else ci[endpoint]['kind']=='unresolved_reference' and ci[endpoint]['name']==str(key))
 if not okay:badrels.append(r['id'])
check('59_source_relationships_exact_endpoints',len(rr)==59 and relrows.max_row-1==59 and Counter(matchedrows)==Counter(range(2,relrows.max_row+1)) and not badrels,dict(edges=len(rr),bad_relationship_ids=badrels))
check('configuration_is_not_live_proof',all(r['evidence_level']=='configured' and r['observed_condition']=='unknown' and r['observed_at'] is None for r in rr) and all(c['observed_condition']=='unknown' and c['observed_at'] is None and c['evidence_level']!='tested' for c in cc),'Inspected every graph node and edge proof/condition field')
bind=[t for t in tables if t['sheet']=='S04-21 Worker Binding Fields'][0];bt=json.loads((ROOT/bind['path']).read_text())
check('201_binding_fields_do_not_create_false_edges',w['S04-21 Worker Binding Fields'].max_row-1==201 and len([r for r in bt['rows'] if r['row']>=2])==201 and all(ei[e]['sheet']!='S04-21 Worker Binding Fields' for r in rr for e in r['evidence_ids']),'201 original rows preserved; none supplies any factual relationship edge')
view=json.loads((ROOT/'views/overview.json').read_text())
member_ids={x['component_id'] for x in view['memberships']}
check('view_contains_all_components_and_navigation_parents',member_ids==set(ci) and all(not x['parent_id'] or x['parent_id'] in member_ids for x in view['memberships']),dict(view_components=len(member_ids),canonical_components=len(ci),missing_component_ids=sorted(set(ci)-member_ids)))
check('view_memberships_unique',len({(x['component_id'],x['parent_id']) for x in view['memberships']})==len(view['memberships']),dict(count=len(view['memberships'])))
map_members=[]
for x in maprows:
 match=[v for v in view['memberships'] if v['order']==x['map_row'] and v['parent_id'] is None and ci[v['component_id']]['name']==x['name']]
 if len(match)==1:map_members.append(match[0]['component_id'])
check('198_distinct_original_overview_entries',len(map_members)==len(set(map_members))==198,dict(matched_original_map_rows=len(map_members),unique_map_components=len(set(map_members))))
check('presentation_membership_separate',view['membership_is_runtime_relationship'] is False and set(view['relationship_ids'])=={r['id'] for r in rr},dict(memberships=len(view['memberships']),factual_relationships=len(rr)))
badpages=[]
for c in cc:
 p=ROOT/'components'/f"{c['id']}.md"
 if not p.is_file():badpages.append(c['id']);continue
 text=p.read_text()
 if c['id'] not in text or c['name'] not in text or (c['purpose'] and c['purpose'] not in text) or (c['boundary'] and c['boundary'] not in text) or c['observed_condition'] not in text or any(e not in text for e in c['evidence_ids']) or any(r not in text for r in c['source_record_ids']) or any(r['id'] not in text for r in rr if c['id'] in (r['source_id'],r['target_id'])):badpages.append(c['id'])
check('all_component_pages_reopened',not badpages,dict(count=len(cc),bad_page_ids=badpages))
page_ids={p.stem for p in (ROOT/'components').glob('*.md')}
check('no_orphan_generated_component_pages',page_ids==set(ci),dict(pages=len(page_ids),canonical=len(ci),orphan_ids=sorted(page_ids-set(ci))))
companion=openpyxl.load_workbook(ROOT/'views/Alex_System_Foundation.xlsx')
norm=lambda v: '' if v is None else v
validviews=True
bylayer={x['id']:x['name'] for x in layers}
expected_top={c['id']:[bylayer[c['layer_id']],c['name'],c['kind'],c['category'],c['evidence_level'],c['observed_condition'],c['purpose'],c['boundary'],c['id']] for c in cc}
actual_top={row[-1].value:[x.value for x in row] for row in list(companion['TOP DOWN MAP'])[1:]}
validviews &= set(actual_top)==set(expected_top) and all([norm(x) for x in actual_top[k]]==[norm(x) for x in v] for k,v in expected_top.items())
expected_rel={r['id']:[r['source_id'],r['type'],r['target_id'],r['label'],r['evidence_level'],r['observed_condition'],r['resolution'],r['id']] for r in rr}
actual_rel={row[-1].value:[x.value for x in row] for row in list(companion['RELATIONSHIPS'])[1:]}
validviews &= set(actual_rel)==set(expected_rel) and all([norm(x) for x in actual_rel[k]]==[norm(x) for x in v] for k,v in expected_rel.items())
expected_cov={r['record_id']:[r['record_id'],r['record_type'],r['component_id'],r['source']['sheet'],r['source']['row']] for r in coverage}
actual_cov={row[0].value:[x.value for x in row] for row in list(companion['COVERAGE'])[1:]}
validviews &= set(actual_cov)==set(expected_cov) and all(actual_cov[k]==v for k,v in expected_cov.items())
check('companion_workbook_matches_canonical_graph_and_coverage',validviews,dict(components=len(actual_top),relationships=len(actual_rel),coverage=len(actual_cov)))
check('companion_workbook_no_formulas_or_errors',not any(c.data_type in ('f','e') for sheet in companion for row in sheet for c in row),'Reopened generated XLSX and inspected every cell type')
summary=dict(scope='Canonical source accounting; no services/UI/app writes',results=results,ui_status='UNVERIFIED: Cursor application/UI build remains pending',checked_sha256={str(p.relative_to(ROOT)):sha(p) for p in sorted((ROOT/'data').glob('*.json'))})
(ROOT/'reports/independent-foundation-graph.json').write_text(json.dumps(summary,indent=2)+'\n')
print(json.dumps(summary,indent=2));sys.exit(1 if any(r['status']=='FAILED' for r in results) else 0)
